"""Build the ZitFlow training guide (Word + PDF) from the Markdown chapters.

Usage (from docs/formation):
    build/.venv/Scripts/python build/build_guide.py [--no-pdf]
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
REPO = ROOT.parent.parent
DIST = ROOT / "dist"
OUTPUT_NAME = "Guide-formation-ZitFlow"
CHAPTERS = [
    "00-introduction.md",
    "01-interface.md",
    "02-reception.md",
    "annexe-a-revue-layouts.md",
    "annexe-b-revue-reception.md",
]
LOGO = REPO / "src" / "assets" / "logos" / "zitflow-lockup.png"

GREEN = RGBColor(0x2E, 0x7D, 0x32)
DARK = RGBColor(0x1F, 0x2A, 0x24)
GREY = RGBColor(0x6B, 0x72, 0x80)
FONT = "Calibri"
HEADER_FILL = "2E7D32"
ZEBRA_FILL = "F1F8F2"
NOTE_FILLS = {"astuce": "E8F5E9", "attention": "FDECEA", "remarque": "EEF2F7"}
PLACEHOLDER_FILL = "F3F4F6"

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


# ---------------------------------------------------------------- helpers

def shade(element, fill: str) -> None:
    pr = element.get_or_add_tcPr() if element.tag.endswith("}tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def para_border(paragraph, color: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    borders.append(left)
    ppr.append(borders)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)
    run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    if placeholder:
        paragraph.add_run(placeholder)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def add_inline(paragraph, text: str, bold: bool = False, color: RGBColor | None = None) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.bold = bold
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            run.bold = bold
        else:
            run = paragraph.add_run(part)
            run.bold = bold
        if color is not None:
            run.font.color.rgb = color


# ---------------------------------------------------------------- styles

def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    sizes = {1: 20, 2: 15, 3: 12.5}
    for level, size in sizes.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = GREEN if level < 3 else DARK
        style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.space_before = Pt(18 if level == 1 else 12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number", "List Bullet 2"):
        doc.styles[name].font.name = FONT
        doc.styles[name].font.size = Pt(10.5)


def setup_section(section) -> None:
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)


def setup_header_footer(section) -> None:
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.text = ""
    run = header.add_run("ZitFlow — Guide de formation")
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Page ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY
    add_field(footer, "PAGE", "1")
    run = footer.add_run(" / ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY
    add_field(footer, "NUMPAGES", "1")


# ---------------------------------------------------------------- blocks

def cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Cm(8))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    run = p.add_run("Guide de formation")
    run.font.size = Pt(30)
    run.bold = True
    run.font.color.rgb = GREEN
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Fonctionnement complet de l'application")
    run.font.size = Pt(15)
    run.font.color.rgb = DARK
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    months = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet",
              "août", "septembre", "octobre", "novembre", "décembre"]
    today = dt.date.today()
    run = p.add_run(f"Version du {today.day} {months[today.month - 1]} {today.year} — usage interne équipe de formation")
    run.font.size = Pt(10)
    run.font.color.rgb = GREY
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def toc(doc: Document) -> None:
    h = doc.add_paragraph()
    h.paragraph_format.space_after = Pt(12)
    run = h.add_run("Sommaire")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = GREEN
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-2" \\h \\z \\u',
              "Le sommaire se met à jour à l'ouverture (ou clic droit › Mettre à jour les champs).")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def image(doc: Document, base: Path, alt: str, target: str) -> None:
    path = (base / target).resolve()
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(15.5))
    else:
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        shade(cell._tc, PLACEHOLDER_FILL)
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(28)
        cp.paragraph_format.space_after = Pt(28)
        run = cp.add_run(f"Capture à insérer : {Path(target).name}")
        run.italic = True
        run.font.color.rgb = GREY
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(alt)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GREY


def note(doc: Document, text: str) -> None:
    kind = "remarque"
    match = re.match(r"\*\*([^*:]+)\s*:\*\*", text)
    if match:
        key = match.group(1).strip().lower()
        kind = next((k for k in NOTE_FILLS if key.startswith(k)), "remarque")
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    shade(p._p, NOTE_FILLS[kind])
    para_border(p, {"astuce": "2E7D32", "attention": "C62828", "remarque": "5B6B82"}[kind])
    add_inline(p, text)


def table(doc: Document, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(header):
        cell = t.rows[0].cells[i]
        shade(cell._tc, HEADER_FILL)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        add_inline(cell.paragraphs[0], text, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for r, row in enumerate(body, start=1):
        for i in range(len(header)):
            cell = t.rows[r].cells[i]
            if r % 2 == 0:
                shade(cell._tc, ZEBRA_FILL)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(2)
            add_inline(cp, row[i] if i < len(row) else "")
            for run in cp.runs:
                run.font.size = Pt(9.5)
    tr = t.rows[0]._tr
    trpr = tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def code_block(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    shade(p._p, PLACEHOLDER_FILL)
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


# ---------------------------------------------------------------- markdown

def render(doc: Document, path: Path, first: bool) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    if not first:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            code_block(doc, block)
            i += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            doc.add_paragraph(heading.group(2), style=f"Heading {len(heading.group(1))}")
            i += 1
            continue
        img = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if img:
            image(doc, path.parent, img.group(1), img.group(2))
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\|[\s:|-]+\|$", lines[i].strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            table(doc, rows)
            continue
        quote = re.match(r"^(?:-\s+)?>\s?(.*)$", stripped)
        if quote:
            text = quote.group(1)
            i += 1
            while i < len(lines) and lines[i].strip().startswith(">"):
                text += " " + lines[i].strip()[1:].strip()
                i += 1
            note(doc, text)
            continue
        bullet = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if bullet:
            level = 2 if len(bullet.group(1)) >= 2 else 1
            p = doc.add_paragraph(style="List Bullet 2" if level == 2 else "List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, bullet.group(2))
            i += 1
            continue
        numbered = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if numbered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(f"{numbered.group(1)}.\t")
            run.bold = True
            run.font.color.rgb = GREEN
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            add_inline(p, numbered.group(2))
            i += 1
            continue
        text = stripped
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|!\[|\||>|[-*]\s|\d+\.\s|```)", lines[i].strip()):
            text += " " + lines[i].strip()
            i += 1
        add_inline(doc.add_paragraph(), text)


def export_pdf(docx: Path, pdf: Path) -> None:
    script = f"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
try {{
  $doc = $word.Documents.Open('{docx}')
  $doc.TablesOfContents | ForEach-Object {{ $_.Update() }}
  $doc.Fields.Update() | Out-Null
  $doc.Save()
  $doc.ExportAsFixedFormat('{pdf}', 17)
  $doc.Close()
}} finally {{ $word.Quit() }}
"""
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--no-pdf", action="store_true")
    args = parser.parse_args()

    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    setup_section(section)
    setup_header_footer(section)
    cover(doc)
    toc(doc)
    for index, name in enumerate(CHAPTERS):
        render(doc, ROOT / name, first=index == 0)

    DIST.mkdir(exist_ok=True)
    docx = DIST / f"{OUTPUT_NAME}.docx"
    doc.save(docx)
    print(f"Word : {docx}")
    if not args.no_pdf:
        pdf = DIST / f"{OUTPUT_NAME}.pdf"
        export_pdf(docx, pdf)
        print(f"PDF  : {pdf}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
